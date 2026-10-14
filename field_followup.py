"""SA2 Risk / Field Follow-up Log (one merged page) - NOC tracking of field teams for
NSA1-4 / SA1-4 tickets (SA2 rows carry extra Risk fields), following the Escalation & Tracking Flow (notify technician <=10 min,
notify NOP <=15 min, chase Depart at T+30, escalate Supervisor NZ -> Manager
Zone -> Region Manager from T+1 h, executive-room updates, MateLine milestones,
close).

Storage: its OWN Google Sheet (FIELD_FOLLOW_SHEET_ID) - deliberately NOT the
Ticket Planing sheet, which the app wipes and rewrites on every build. Nothing
here is ever deleted.

    Tickets  - one row per ticket (latest state: step times/by/notes, MateLine
               milestones, executive updates, level reached, closed time)
    Log      - append-only audit trail, one row per button press
    Editors  - Email | Name | Role (editor / admin). Edit it in the sheet, no
               redeploy needed (cached ~60 s). People not listed can only view.

The tabs are created automatically on first use. Editors is seeded once (only
when it is empty) from DEFAULT_EDITORS below.

Permissions are enforced HERE (server side), not just by greying out buttons.
Gunicorn runs a single worker (see oncall.py), so the in-memory caches below
are shared by every request; writes update the cache in place.
"""

import json
import logging
import os
import re
import threading
import time
from datetime import datetime, timedelta
from io import BytesIO

import pytz

import sa2_risk  # option lists (INOC names, Topo Ring, channel) + the old SA2 Risk sheet for the one-time import

log = logging.getLogger(__name__)

TZ = pytz.timezone("Asia/Bangkok")
SHEET_ID = os.environ.get("FIELD_FOLLOW_SHEET_ID", "1L2rOsYVjakywybRZq31TKiV6x7Ri7yhyOTGrCG63yK0")
TICKETS_SHEET = "Tickets"
LOG_SHEET = "Log"
EDITORS_SHEET = "Editors"

SLA_VALUES = ("NSA1", "NSA2", "NSA3", "NSA4", "SA1", "SA2", "SA3", "SA4")

# step key -> (column prefix, escalation level reached when done, due minutes after receipt or None)
STEPS = {
    "s1": ("S1_Notify_Tech", 0, 10),
    "s2": ("S2_Notify_NOP", 1, 15),
    "s3": ("S3_Chase_Depart", 1, 30),
    "s4a": ("S4_Supervisor_NZ", 2, 60),
    "s4b": ("S4_Manager_Zone", 3, None),
    "s4c": ("S4_Region_Manager", 4, None),
}
STEP_ORDER = ["s1", "s2", "s3", "s4a", "s4b", "s4c"]

# MateLine milestones: key -> single-time column ("update" is repeatable -> count + list)
MATE_KEYS = ["depart", "arrive", "update", "complete", "closed", "hold"]
_MATE_COL = {
    "depart": "Mate_Depart_At", "arrive": "Mate_Arrive_At", "complete": "Mate_Complete_At",
    "closed": "Mate_Closed_At", "hold": "Mate_Hold_At",
}

TICKET_HEADER = (
    ["Ticket_ID", "SLA", "Subject", "Province", "Impacted_Sites", "Received_At", "Status"]
    + [f"{STEPS[k][0]}_{suffix}" for k in STEP_ORDER for suffix in ("At", "By", "Note", "Skip")]
    + ["Exec_Updates", "Mate_Depart_At", "Mate_Arrive_At", "Mate_Update_Count", "Mate_Update_Times",
       "Mate_Complete_At", "Mate_Closed_At", "Mate_Hold_At", "Max_Level",
       "Closed_At", "Closed_By", "Created_By", "Updated_At", "Updated_By",
       "Deleted_At", "Deleted_By", "Delete_Reason",
       # SA2 Risk fields (only filled on SLA = SA2 rows) + Import_Key (set only on rows brought in from the old SA2_Risk sheet)
       "INOC_Name", "Ticket_Create_At", "Topo_Ring", "Channel", "Import_Key"]
)
LOG_HEADER = ["Time", "Ticket_ID", "Step", "Action", "By", "Note"]
EDITORS_HEADER = ["Email", "Name", "Role"]

# Seeded into the Editors tab ONLY when that tab is empty. After that the sheet is the source of truth.
DEFAULT_EDITORS = [
    ("saridphong_n@bbtec.co.th", "", "admin"),
    ("ponrawe.s@rfplusthailand.com", "", "editor"),
    ("watcharaporn.s@rfplusthailand.com", "", "editor"),
    ("charinthip.k@rfplusthailand.com", "", "editor"),
    ("wanitcha.n@rfplusthailand.com", "", "editor"),
    ("kitthipong.j@rfplusthailand.com", "", "editor"),
    ("rapeepath.b@rfplusthailand.com", "", "editor"),
    ("thanachot.c@rfplusthailand.com", "", "editor"),
    ("naphat.s@rfplusthailand.com", "", "editor"),
    ("matas.n@rfplusthailand.com", "", "editor"),
    ("charansin.c@rfplusthailand.com", "", "editor"),
    ("natchawan.b@rfplusthailand.com", "", "editor"),
]

FALLBACK_PROVINCES = [
    "เชียงใหม่", "เชียงราย", "ลำปาง", "ลำพูน", "แม่ฮ่องสอน", "น่าน", "พะเยา", "แพร่",
    "อุตรดิตถ์", "ตาก", "กำแพงเพชร", "สุโขทัย", "พิษณุโลก", "เพชรบูรณ์", "พิจิตร",
]

_TICKETS_TTL = 120
_EDITORS_TTL = 60
_lock = threading.RLock()
_state = {"header": None, "raw": [], "index": {}, "ts": 0, "ws": None, "ready": False}
_editors = {"rows": None, "ts": 0}
_ensure_done = {"ok": False}


# ---------------------------------------------------------------- time helpers
def _now():
    return datetime.now(TZ)


def _fmt(dt):
    return dt.strftime("%Y-%m-%d %H:%M:%S")


def _now_str():
    return _fmt(_now())


def _to_ms(s):
    """'YYYY-MM-DD HH:MM:SS' (Bangkok) -> epoch ms, or None."""
    if not s:
        return None
    try:
        return int(TZ.localize(datetime.strptime(str(s).strip(), "%Y-%m-%d %H:%M:%S")).timestamp() * 1000)
    except ValueError:
        return None


def _clean(s, limit=500):
    return str(s or "").strip()[:limit]


# ---------------------------------------------------------------- sheet setup
def _spreadsheet(gs):
    return gs.open_by_key(SHEET_ID)


def _get_or_create(sh, title, header, rows=1000):
    try:
        ws = sh.worksheet(title)
        created = False
    except Exception:
        ws = sh.add_worksheet(title=title, rows=rows, cols=max(len(header) + 2, 8))
        created = True
    return ws, created


def _ensure(gs):
    """Creates the three tabs (and the header rows) if they do not exist yet. Cheap after the first call."""
    with _lock:
        if _ensure_done["ok"]:
            return
        sh = _spreadsheet(gs)

        ws, created = _get_or_create(sh, TICKETS_SHEET, TICKET_HEADER, rows=2000)
        first = ws.row_values(1)
        if not first:
            ws.update("A1", [TICKET_HEADER], value_input_option="RAW")
            try:
                ws.freeze(rows=1, cols=1)
            except Exception:
                pass
        else:
            missing = [h for h in TICKET_HEADER if h not in first]
            if missing:  # forward-compatible: only ever appends columns
                if ws.col_count < len(first) + len(missing):
                    ws.add_cols(len(first) + len(missing) - ws.col_count)
                from gspread.utils import rowcol_to_a1
                ws.update(rowcol_to_a1(1, len(first) + 1), [missing], value_input_option="RAW")

        lws, _ = _get_or_create(sh, LOG_SHEET, LOG_HEADER, rows=5000)
        if not lws.row_values(1):
            lws.update("A1", [LOG_HEADER], value_input_option="RAW")
            try:
                lws.freeze(rows=1)
            except Exception:
                pass

        ews, _ = _get_or_create(sh, EDITORS_SHEET, EDITORS_HEADER, rows=200)
        if not ews.row_values(1):
            ews.update("A1", [EDITORS_HEADER] + [list(r) for r in DEFAULT_EDITORS], value_input_option="RAW")

        _ensure_done["ok"] = True


# ---------------------------------------------------------------- editors / roles
def _load_editors(gs):
    now = time.monotonic()
    with _lock:
        if _editors["rows"] is not None and (now - _editors["ts"]) < _EDITORS_TTL:
            return _editors["rows"]
    _ensure(gs)
    ws = _spreadsheet(gs).worksheet(EDITORS_SHEET)
    values = ws.get_all_values()
    rows = {}
    for line in values[1:]:
        if not line or not (line[0] or "").strip():
            continue
        email = line[0].strip().lower()
        role = (line[2].strip().lower() if len(line) > 2 else "") or "editor"
        if role not in ("editor", "admin"):
            role = "editor"
        rows[email] = {"email": email, "name": line[1].strip() if len(line) > 1 else "", "role": role}
    with _lock:
        _editors["rows"] = rows
        _editors["ts"] = now
    return rows


def get_role(gs, email):
    """'admin' / 'editor' / None (view only)."""
    rows = _load_editors(gs)
    r = rows.get((email or "").strip().lower())
    return r["role"] if r else None


def _require_editor(gs, user):
    role = get_role(gs, user["email"])
    if role not in ("editor", "admin"):
        raise PermissionError("บัญชีนี้ไม่มีสิทธิ์แก้ไข (ดูได้อย่างเดียว)")
    return role


# ---------------------------------------------------------------- tickets cache
def _load(gs, force=False):
    now = time.monotonic()
    with _lock:
        if not force and _state["ready"] and (now - _state["ts"]) < _TICKETS_TTL:
            return
    _ensure(gs)
    ws = _spreadsheet(gs).worksheet(TICKETS_SHEET)
    values = ws.get_all_values()
    header = values[0] if values else list(TICKET_HEADER)
    raw, index = [], {}
    for i, line in enumerate(values[1:], start=2):
        padded = list(line) + [""] * (len(header) - len(line))
        raw.append(padded)
        tid = (padded[header.index("Ticket_ID")] if "Ticket_ID" in header else "").strip()
        if tid:
            index[tid] = len(raw) - 1  # position in raw; sheet row = position + 2
    with _lock:
        _state.update(header=header, raw=raw, index=index, ts=now, ws=ws, ready=True)


def _col(header, name):
    return header.index(name)


def _row_to_ticket(header, row):
    def g(name):
        return row[header.index(name)].strip() if name in header and header.index(name) < len(row) else ""

    steps = {}
    for k in STEP_ORDER:
        p = STEPS[k][0]
        at = g(f"{p}_At")
        if at:
            steps[k] = {"at": at, "by": g(f"{p}_By"), "note": g(f"{p}_Note"), "skip": g(f"{p}_Skip").upper() == "Y"}
    try:
        upd = json.loads(g("Exec_Updates") or "[]")
        if not isinstance(upd, list):
            upd = []
    except ValueError:
        upd = []
    mate = {k: g(c) for k, c in _MATE_COL.items()}
    mate["update"] = [x.strip() for x in g("Mate_Update_Times").split(",") if x.strip()]
    try:
        sites = int(float(g("Impacted_Sites") or 0))
    except ValueError:
        sites = 0
    return {
        "id": g("Ticket_ID"), "sla": g("SLA"), "subj": g("Subject"), "prov": g("Province"), "sites": sites,
        "recv": g("Received_At"), "status": g("Status") or "Open", "steps": steps, "upd": upd, "mate": mate,
        "max_level": g("Max_Level"), "closed_at": g("Closed_At"), "closed_by": g("Closed_By"),
        "created_by": g("Created_By"), "updated_at": g("Updated_At"), "updated_by": g("Updated_By"),
        "deleted_at": g("Deleted_At"), "deleted_by": g("Deleted_By"), "delete_reason": g("Delete_Reason"),
        "inoc": g("INOC_Name"), "create": g("Ticket_Create_At"), "topo": g("Topo_Ring"), "ch": g("Channel"),
        "import_key": g("Import_Key"),
    }


def _max_level(t):
    lvl = -1
    for k, v in t["steps"].items():
        if not v.get("skip") and STEPS[k][1] > lvl:
            lvl = STEPS[k][1]
    return lvl


def _ticket_to_row(header, t, base=None):
    row = list(base) if base else [""] * len(header)
    row += [""] * (len(header) - len(row))

    def s(name, val):
        if name in header:
            row[header.index(name)] = "" if val is None else val

    s("Ticket_ID", t["id"]); s("SLA", t["sla"]); s("Subject", t["subj"]); s("Province", t["prov"])
    s("Impacted_Sites", t["sites"]); s("Received_At", t["recv"]); s("Status", t["status"])
    for k in STEP_ORDER:
        p = STEPS[k][0]
        v = t["steps"].get(k)
        s(f"{p}_At", v["at"] if v else ""); s(f"{p}_By", v["by"] if v else "")
        s(f"{p}_Note", v["note"] if v else ""); s(f"{p}_Skip", "Y" if v and v.get("skip") else "")
    s("Exec_Updates", json.dumps(t["upd"], ensure_ascii=False) if t["upd"] else "")
    for k, c in _MATE_COL.items():
        s(c, t["mate"].get(k, ""))
    s("Mate_Update_Count", len(t["mate"]["update"]) or "")
    s("Mate_Update_Times", ",".join(t["mate"]["update"]))
    lvl = _max_level(t)
    s("Max_Level", "" if lvl < 0 else lvl)
    s("Closed_At", t["closed_at"]); s("Closed_By", t["closed_by"]); s("Created_By", t["created_by"])
    s("Updated_At", t["updated_at"]); s("Updated_By", t["updated_by"])
    s("Deleted_At", t.get("deleted_at", "")); s("Deleted_By", t.get("deleted_by", "")); s("Delete_Reason", t.get("delete_reason", ""))
    s("INOC_Name", t.get("inoc", "")); s("Ticket_Create_At", t.get("create", "")); s("Topo_Ring", t.get("topo", ""))
    s("Channel", t.get("ch", "")); s("Import_Key", t.get("import_key", ""))
    return ["" if v is None else str(v) for v in row]  # strings only, exactly what a sheet read returns


def _public(t):
    return {
        "id": t["id"], "sla": t["sla"], "subj": t["subj"], "prov": t["prov"], "sites": t["sites"],
        "recv": _to_ms(t["recv"]), "status": t["status"],
        "steps": {k: {"t": _to_ms(v["at"]), "by": v["by"], "note": v["note"], "skip": bool(v.get("skip"))}
                  for k, v in t["steps"].items()},
        "mate": {
            "depart": _to_ms(t["mate"].get("depart")), "arrive": _to_ms(t["mate"].get("arrive")),
            "complete": _to_ms(t["mate"].get("complete")), "closed": _to_ms(t["mate"].get("closed")),
            "hold": _to_ms(t["mate"].get("hold")),
            "update": [m for m in (_to_ms(x) for x in t["mate"]["update"]) if m],
        },
        "upd": [{"t": _to_ms(u.get("at")), "by": u.get("by", ""), "note": u.get("note", "")} for u in t["upd"]],
        "closed": _to_ms(t["closed_at"]) if t["status"] == "Closed" else None,
        "closed_by": t["closed_by"], "created_by": t["created_by"],
        "deleted": _to_ms(t.get("deleted_at")) if t["status"] == "Deleted" else None,
        "deleted_by": t.get("deleted_by", ""), "delete_reason": t.get("delete_reason", ""),
        "level": _max_level(t),
        "inoc": t.get("inoc", ""), "topo": t.get("topo", "") or ("Send done" if t["sla"] == "SA2" else ""),
        "ch": t.get("ch", ""), "create": _to_ms(t.get("create")) or _to_ms(t["recv"]),
        "imported": bool(t.get("import_key")),
    }


def _find(tid):
    pos = _state["index"].get(tid)
    if pos is None:
        raise ValueError(f"ไม่พบ Ticket {tid}")
    return pos


def _save(gs, t, pos=None):
    """Writes one ticket row (new when pos is None) and updates the cache in place."""
    header = _state["header"]
    ws = _state["ws"]
    if pos is None:
        row = _ticket_to_row(header, t)
        resp = ws.append_row(row, value_input_option="RAW", table_range="A1")
        m = re.search(r"![A-Z]+(\d+):", (resp or {}).get("updates", {}).get("updatedRange", ""))
        sheet_row = int(m.group(1)) if m else len(_state["raw"]) + 2
        # keep raw positions aligned with sheet rows (blank rows, if any, become empty raw rows)
        while len(_state["raw"]) < sheet_row - 2:
            _state["raw"].append([""] * len(header))
        _state["raw"].append(row)
        _state["index"][t["id"]] = len(_state["raw"]) - 1
    else:
        row = _ticket_to_row(header, t, _state["raw"][pos])
        from gspread.utils import rowcol_to_a1
        ws.update(f"{rowcol_to_a1(pos + 2, 1)}", [row], value_input_option="RAW")
        _state["raw"][pos] = row


def _log(gs, tid, step, action, by, note=""):
    try:
        ws = _spreadsheet(gs).worksheet(LOG_SHEET)
        ws.append_row([_now_str(), tid, step, action, by, note], value_input_option="RAW", table_range="A1")
    except Exception:
        log.exception("field-followup log append failed (ticket %s, %s)", tid, action)


def _touch(t, user):
    t["updated_at"] = _now_str()
    t["updated_by"] = user["email"]


def _open_ticket(tid):
    t = _row_to_ticket(_state["header"], _state["raw"][_find(tid)])
    return t


# ---------------------------------------------------------------- reads
def board(gs, user):
    _load(gs)
    cutoff_ms = (_now() - timedelta(hours=24)).timestamp() * 1000
    out, closed_recent = [], []
    with _lock:
        for row in _state["raw"]:
            t = _row_to_ticket(_state["header"], row)
            if not t["id"] or t["status"] == "Deleted":
                continue
            if t["status"] != "Closed":
                out.append(_public(t))          # the board shows OPEN tickets only; closed ones live in History
            else:
                ms = _to_ms(t["recv"])
                if ms and ms >= cutoff_ms:       # kept only so the "last 24 h" summary card can count them
                    closed_recent.append(_public(t))
    return {
        "me": {"email": user["email"], "name": user.get("name") or "", "role": get_role(gs, user["email"])},
        "now": int(_now().timestamp() * 1000),
        "tickets": out,
        "closed_recent": closed_recent,
        "provinces": _provinces(gs),
        "options": {"inoc": list(sa2_risk.INOC_NAMES), "topo": list(sa2_risk.TOPO_RING_OPTIONS),
                    "channel": list(sa2_risk.CHANNEL_OPTIONS)},
    }


def history(gs, user, q="", date_from="", date_to="", sla="", limit=200, show_deleted=False):
    _load(gs)
    show_deleted = bool(show_deleted) and get_role(gs, user["email"]) == "admin"  # only admins can see deleted tickets
    q = (q or "").strip().lower()
    res = []
    with _lock:
        for row in _state["raw"]:
            t = _row_to_ticket(_state["header"], row)
            if not t["id"]:
                continue
            if t["status"] == "Deleted" and not show_deleted:
                continue
            if sla and t["sla"] != sla:
                continue
            if q and q not in f"{t['id']} {t['subj']} {t['prov']}".lower():
                continue
            d = (t["recv"] or "")[:10]
            if date_from and d < date_from:
                continue
            if date_to and d > date_to:
                continue
            res.append(t)
    res.sort(key=lambda t: t["recv"], reverse=True)
    return {"tickets": [_public(t) for t in res[:limit]], "total": len(res)}


def export_xlsx(gs, user, q="", date_from="", date_to="", sla=""):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    data = history(gs, user, q, date_from, date_to, sla, limit=100000)
    wb = Workbook()
    ws = wb.active
    ws.title = "Field Follow-up"
    ws.append(TICKET_HEADER)
    with _lock:
        by_id = {}
        for row in _state["raw"]:
            tid = (row[_state["header"].index("Ticket_ID")] or "").strip()
            if tid:
                by_id[tid] = row
        for t in data["tickets"]:
            row = by_id.get(t["id"], [])
            ws.append([row[_state["header"].index(h)] if h in _state["header"] and _state["header"].index(h) < len(row) else ""
                       for h in TICKET_HEADER])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D4ED8")
    ws.freeze_panes = "B2"
    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio


# ---------------------------------------------------------------- writes
def _hhmm_to_dt(hhmm, now):
    m = re.match(r"^(\d{1,2}):(\d{2})$", (hhmm or "").strip())
    if not m:
        return None
    cand = now.replace(hour=int(m.group(1)) % 24, minute=int(m.group(2)), second=0, microsecond=0)
    if cand > now + timedelta(minutes=5):  # typed a time "later today" -> it was yesterday
        cand -= timedelta(days=1)
    return cand


def _check_sa2_fields(inoc, topo, ch):
    inoc, topo, ch = _clean(inoc, 60), _clean(topo, 30), _clean(ch, 30)
    if not inoc:
        raise ValueError("SA2 ต้องเลือก INOC Name")
    if topo not in sa2_risk.TOPO_RING_OPTIONS:
        raise ValueError("Topo Ring ไม่ถูกต้อง")
    if ch not in sa2_risk.CHANNEL_OPTIONS:
        raise ValueError("ช่องทางแจ้งทีมไม่ถูกต้อง")
    return inoc, topo, ch


def create_ticket(gs, user, tid, sla, subj, prov, sites, recv_hhmm="", inoc="", create_hhmm="", topo="", ch=""):
    _require_editor(gs, user)
    tid, subj, prov = _clean(tid, 60), _clean(subj, 400), _clean(prov, 80)
    if not tid or not subj or not prov:
        raise ValueError("กรุณากรอกเลข Ticket, Subject และจังหวัด")
    if sla not in SLA_VALUES:
        raise ValueError("SLA ต้องเป็น NSA1-4 หรือ SA1-4")
    try:
        sites = int(sites or 0)
    except (TypeError, ValueError):
        raise ValueError("จำนวนไซต์ต้องเป็นตัวเลข")
    if sites < 0 or sites > 5000:
        raise ValueError("จำนวนไซต์ไม่ถูกต้อง")
    if sla == "SA2" and sites < 1:
        raise ValueError("SA2 ต้องระบุจำนวนไซต์ที่กระทบ (อย่างน้อย 1)")

    now = _now()
    recv = _hhmm_to_dt(recv_hhmm, now) or now
    created = recv
    if sla == "SA2":
        inoc, topo, ch = _check_sa2_fields(inoc, topo or "Send done", ch or "Tel")
        created = _hhmm_to_dt(create_hhmm, now) or recv
    else:
        inoc = topo = ch = ""
        created = None

    with _lock:
        _load(gs, force=True)
        if tid in _state["index"]:
            old = _row_to_ticket(_state["header"], _state["raw"][_state["index"][tid]])
            raise ValueError("เลข Ticket นี้เคยถูกลบไปแล้ว (admin กู้คืนได้ที่ History)" if old["status"] == "Deleted" else "มี Ticket นี้ในระบบแล้ว")
        t = {"id": tid, "sla": sla, "subj": subj, "prov": prov, "sites": sites, "recv": _fmt(recv),
             "status": "Open", "steps": {}, "upd": [], "mate": {k: "" for k in _MATE_COL} | {"update": []},
             "max_level": "", "closed_at": "", "closed_by": "", "created_by": user["email"],
             "updated_at": "", "updated_by": "", "inoc": inoc, "topo": topo, "ch": ch,
             "create": _fmt(created) if created else "", "import_key": ""}
        _touch(t, user)
        _save(gs, t)
    _log(gs, tid, "-", "CREATE", user["email"],
         f"{sla} · {prov} · {sites} sites · {subj}" + (f" · {inoc} · Topo {topo} · {ch}" if sla == "SA2" else ""))
    return _public(t)


def _mutate(gs, user, tid, fn, allow_closed=False):
    _require_editor(gs, user)
    with _lock:
        _load(gs)
        pos = _find(tid)
        t = _row_to_ticket(_state["header"], _state["raw"][pos])
        if t["status"] == "Closed" and not allow_closed:
            raise ValueError("Ticket นี้ปิดแล้ว (admin เปิดใหม่ได้)")
        fn(t)
        _touch(t, user)
        _save(gs, t, pos)
    return _public(t)


def step_action(gs, user, tid, key, action, note=""):
    if key not in STEPS:
        raise ValueError("ขั้นตอนไม่ถูกต้อง")
    if action not in ("done", "skip"):
        raise ValueError("action ต้องเป็น done หรือ skip")
    note = _clean(note)

    def fn(t):
        if key in t["steps"]:
            raise ValueError("ขั้นตอนนี้บันทึกไปแล้ว")
        t["steps"][key] = {"at": _now_str(), "by": user["email"], "note": note, "skip": action == "skip"}

    res = _mutate(gs, user, tid, fn)
    _log(gs, tid, key, "STEP_DONE" if action == "done" else "STEP_SKIP", user["email"], note)
    return res


def set_remark(gs, user, tid, key, note):
    if key not in STEPS:
        raise ValueError("ขั้นตอนไม่ถูกต้อง")
    note = _clean(note)

    def fn(t):
        if key not in t["steps"]:
            raise ValueError("ขั้นตอนนี้ยังไม่ได้บันทึก")
        t["steps"][key]["note"] = note

    res = _mutate(gs, user, tid, fn, allow_closed=True)
    _log(gs, tid, key, "REMARK", user["email"], note)
    return res


def set_sa2_fields(gs, user, tid, inoc, topo, ch):
    inoc, topo, ch = _check_sa2_fields(inoc, topo, ch)

    def fn(t):
        if t["sla"] != "SA2":
            raise ValueError("ช่อง SA2 Risk ใช้ได้เฉพาะ Ticket ที่ SLA = SA2")
        t["inoc"], t["topo"], t["ch"] = inoc, topo, ch

    res = _mutate(gs, user, tid, fn)
    _log(gs, tid, "sa2", "SA2_FIELDS", user["email"], f"{inoc} · Topo {topo} · {ch}")
    return res


def add_exec_update(gs, user, tid, note):
    note = _clean(note)
    if not note:
        raise ValueError("กรุณาพิมพ์ข้อความ Update")

    def fn(t):
        t["upd"].append({"at": _now_str(), "by": user["email"], "note": note})

    res = _mutate(gs, user, tid, fn)
    _log(gs, tid, "5", "EXEC_UPDATE", user["email"], note)
    return res


def mate_toggle(gs, user, tid, key, on):
    if key not in MATE_KEYS:
        raise ValueError("MateLine milestone ไม่ถูกต้อง")

    def fn(t):
        if key == "update":
            if on:
                t["mate"]["update"].append(_now_str())
            elif t["mate"]["update"]:
                t["mate"]["update"].pop()
        else:
            t["mate"][key] = _now_str() if on else ""

    res = _mutate(gs, user, tid, fn)
    _log(gs, tid, "5.1", ("MATE_" if on else "MATE_OFF_") + key.upper(), user["email"])
    return res


def close_ticket(gs, user, tid):
    def fn(t):
        t["status"] = "Closed"
        t["closed_at"] = _now_str()
        t["closed_by"] = user["email"]

    res = _mutate(gs, user, tid, fn)
    _log(gs, tid, "close", "CLOSE", user["email"])
    return res


def delete_ticket(gs, user, tid, reason=""):
    """Soft delete for tickets opened by mistake: the row STAYS in the sheet (Status = Deleted, who/when/why recorded,
    Log keeps the audit entry) but disappears from the board, summary and normal history. Editors can delete open
    tickets; closed tickets can only be deleted by an admin. Admin can restore."""
    role = _require_editor(gs, user)
    reason = _clean(reason, 200)

    def fn(t):
        if t["status"] == "Deleted":
            raise ValueError("Ticket นี้ถูกลบไปแล้ว")
        if t["status"] == "Closed" and role != "admin":
            raise PermissionError("Ticket ที่ปิดแล้ว ลบได้เฉพาะ admin")
        t["status"] = "Deleted"
        t["deleted_at"] = _now_str()
        t["deleted_by"] = user["email"]
        t["delete_reason"] = reason

    res = _mutate(gs, user, tid, fn, allow_closed=True)
    _log(gs, tid, "delete", "DELETE", user["email"], reason)
    return res


def restore_ticket(gs, user, tid):
    if get_role(gs, user["email"]) != "admin":
        raise PermissionError("เฉพาะ admin เท่านั้นที่กู้คืน Ticket ได้")

    def fn(t):
        if t["status"] != "Deleted":
            raise ValueError("Ticket นี้ไม่ได้ถูกลบ")
        t["status"] = "Closed" if t["closed_at"] else "Open"
        t["deleted_at"] = t["deleted_by"] = t["delete_reason"] = ""

    res = _mutate(gs, user, tid, fn, allow_closed=True)
    _log(gs, tid, "delete", "RESTORE", user["email"])
    return res


def reopen_ticket(gs, user, tid):
    if get_role(gs, user["email"]) != "admin":
        raise PermissionError("เฉพาะ admin เท่านั้นที่เปิด Ticket ที่ปิดแล้วกลับมาได้")

    def fn(t):
        t["status"] = "Open"
        t["closed_at"] = ""
        t["closed_by"] = ""

    res = _mutate(gs, user, tid, fn, allow_closed=True)
    _log(gs, tid, "close", "REOPEN", user["email"])
    return res


# ---------------------------------------------------------------- Oncall lookup (reads the existing Oncall tabs)
def _today():
    return _now().strftime("%Y-%m-%d")


def _provinces(gs):
    names = set()
    try:
        import oncall
        data = oncall.load_oncall_schedule(gs)
        if data:
            names |= {r["province_th"].strip() for r in data["rows"] if r.get("province_th", "").strip()}
        import oncall_escalation
        esc = oncall_escalation.load_escalation_contacts(gs)
        if esc:
            names |= {p["province"].strip() for p in esc["provinces"] if p.get("province", "").strip()}
    except Exception:
        log.exception("field-followup province list failed")
    return sorted(names) if names else list(FALLBACK_PROVINCES)


def oncall_for_province(gs, user, province):
    """Who is Oncall TODAY in `province`, in the order the escalation flow uses them:
    Level 1 field teams (Team ID / Remark / Note / district / Tel), Level 2 Supervisor (+ Supervisor MTE),
    Level 3 Eng Zone. A person counts as Oncall today when their cell for today is "on" (or legacy "always")."""
    province = (province or "").strip()
    today = _today()
    import oncall
    import oncall_escalation

    teams, sched_ok = [], False
    data = oncall.load_oncall_schedule(gs)
    if data is not None:
        sched_ok = today in data["dates"]
        for r in data["rows"]:
            if province.lower() not in ((r.get("province_th") or "").strip().lower(), (r.get("province") or "").strip().lower()):
                continue
            if (r["days"].get(today) or "blank") in ("on", "always"):
                teams.append({"team_id": r.get("team_id1") or r.get("team_id"), "remark": r.get("remark", ""),
                              "note": r.get("note", ""), "district": r.get("district", ""),
                              "tel": r.get("tel", ""), "name": r.get("name", "")})
    teams.sort(key=lambda x: (x["remark"] or "~", x["team_id"] or ""))

    sup_on, eng_on, sup_all, eng_all = [], [], [], []
    esc = oncall_escalation.load_escalation_contacts(gs)
    if esc is not None:
        for p in esc["provinces"]:
            if (p.get("province") or "").strip().lower() != province.lower():
                continue
            for c in p["contacts"]:
                item = {"position": c["position"], "name": c["name"], "type": c["type"],
                        "note": c.get("note", ""), "tel": c["tel"],
                        "on_today": (c["days"].get(today) or "blank") in ("on", "always")}
                bucket_on, bucket_all = (eng_on, eng_all) if c["position"] == "Eng Zone" else (sup_on, sup_all)
                bucket_all.append(item)
                if item["on_today"]:
                    bucket_on.append(item)
    return {
        "date": today, "province": province, "schedule_has_today": sched_ok,
        "level1": teams,
        "level2": sup_on or sup_all, "level2_none_on_today": bool(sup_all) and not sup_on,
        "level3": eng_on or eng_all, "level3_none_on_today": bool(eng_all) and not eng_on,
    }


# ---------------------------------------------------------------- one-time import from the old "SA2 Risk" tab
# The old tab keeps its own sheet (sa2_risk.SA2_RISK_SHEET_ID / "SA2_Risk"). It is only READ here - never changed.
# Every imported row is stamped Import_Key = "SA2R:<old row id>", so pressing the button again can never create a
# second copy: rows whose key is already present are skipped, and so is any ticket number that already exists.
_NOP_TO_PROVINCE = {
    "CMI": "เชียงใหม่", "CRI": "เชียงราย", "LPG": "ลำปาง", "LPN": "ลำพูน", "MHS": "แม่ฮ่องสอน", "NAN": "น่าน",
    "PHE": "แพร่", "PYO": "พะเยา", "KPP": "กำแพงเพชร", "PCB": "เพชรบูรณ์", "PCT": "พิจิตร", "PSN": "พิษณุโลก",
    "SKT": "สุโขทัย", "TAK": "ตาก", "UTR": "อุตรดิตถ์",
}
_DT_FORMATS = ("%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
               "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M")
IMPORT_NOTE = "นำเข้าจาก SA2 Risk เดิม (ไม่มีเวลาบันทึกของขั้นนี้)"
IMPORT_BY = "นำเข้า SA2 Risk"


def _parse_dt(s, server_clock=False):
    """Old-tab date text -> Bangkok datetime, or None. Times a person picked in the form are Bangkok local time;
    updated_at / closed_dt were stamped by the server's clock (UTC on Railway), hence server_clock=True."""
    s = (s or "").strip()
    if not s:
        return None
    for f in _DT_FORMATS:
        try:
            d = datetime.strptime(s, f)
        except ValueError:
            continue
        return pytz.utc.localize(d).astimezone(TZ) if server_clock else TZ.localize(d)
    return None


def _province_from(text):
    text = (text or "").strip()
    m = re.search(r"-([A-Z]{3})-NOP", text)
    if m and m.group(1) in _NOP_TO_PROVINCE:
        return _NOP_TO_PROVINCE[m.group(1)]
    return text or "ไม่ระบุ"


def _import_ticket(r, key, now, user):
    p = lambda k: (r.get(k) or "").strip()
    recv = _parse_dt(p("inoc_open_dt")) or _parse_dt(p("create_dt")) or _parse_dt(p("updated_at"), True) or now
    recv = min(recv, now)
    created = min(_parse_dt(p("create_dt")) or recv, now)
    closed = p("status") == "Closed"
    closed_at = ""
    if closed:
        c = _parse_dt(p("closed_dt"), True) or _parse_dt(p("updated_at"), True) or now
        closed_at = _fmt(max(min(c, now), recv))
    # no per-step times exist in the old tab: mark every step "skipped (imported)" so a pending imported ticket
    # shows "all steps done - waiting to close" instead of a bogus red overdue clock
    steps = {k: {"at": _fmt(recv), "by": IMPORT_BY, "note": IMPORT_NOTE, "skip": True} for k in STEP_ORDER}
    ack = _parse_dt(p("sup_ack_dt"))
    if ack:
        steps["s2"] = {"at": _fmt(ack), "by": p("sup_name") or IMPORT_BY, "skip": False,
                       "note": "นำเข้า: Sup/NOP รับทราบ" + (f" · {p('sup_name')}" if p("sup_name") else "")}
    try:
        sites = max(0, int(float(p("risk_sites") or 0)))
    except ValueError:
        sites = 0
    subj = p("subject") or "(ไม่มี Subject)"
    if p("site_name") and p("site_name") not in subj:
        subj += f" [{p('site_name')}]"
    return {
        "id": p("ticket_id"), "sla": "SA2", "subj": subj[:400], "prov": _province_from(p("province")),
        "sites": sites, "recv": _fmt(recv), "status": "Closed" if closed else "Open", "steps": steps, "upd": [],
        "mate": {k: "" for k in _MATE_COL} | {"update": []}, "max_level": "",
        "closed_at": closed_at, "closed_by": (p("updated_by") or user["email"]) if closed else "",
        "created_by": p("updated_by") or user["email"], "updated_at": _fmt(now), "updated_by": user["email"],
        "inoc": p("inoc_name"), "create": _fmt(created),
        "topo": p("topo_ring") if p("topo_ring") in sa2_risk.TOPO_RING_OPTIONS else "Send done",
        "ch": p("channel") if p("channel") in sa2_risk.CHANNEL_OPTIONS else "",
        "import_key": key,
    }


def import_sa2_risk(gs, user, dry_run=False):
    """Admin only. dry_run=True only counts (nothing is written) so the button can show what WILL happen first."""
    if get_role(gs, user["email"]) != "admin":
        raise PermissionError("เฉพาะ admin เท่านั้นที่นำเข้าข้อมูลจาก SA2 Risk เดิมได้")
    try:
        src = sa2_risk.get_all_rows(gs, use_cache=False)
    except sa2_risk.SA2RiskError as e:
        raise ValueError(str(e))
    now = _now()
    st = {"total": len(src), "new_open": 0, "new_closed": 0, "merged": 0, "already": 0, "dup_in_source": 0,
          "invalid": 0, "dry_run": bool(dry_run)}
    with _lock:
        _load(gs, force=True)
        header = _state["header"]
        kcol = header.index("Import_Key") if "Import_Key" in header else None
        keys = {(row[kcol] or "").strip() for row in _state["raw"] if kcol is not None and kcol < len(row)} - {""}
        seen, new, merges = set(), [], []
        for r in src:
            tid, rid = (r.get("ticket_id") or "").strip(), (r.get("id") or "").strip()
            if not tid or not rid:
                st["invalid"] += 1
                continue
            if tid in seen:
                st["dup_in_source"] += 1   # same ticket typed twice in the old tab: the first row wins
                continue
            seen.add(tid)
            key = "SA2R:" + rid
            if key in keys:
                st["already"] += 1
                continue
            pos = _state["index"].get(tid)
            if pos is not None:            # ticket already in the new page: never overwrite, only fill empty SA2 fields
                t = _row_to_ticket(header, _state["raw"][pos])
                if t["status"] != "Deleted" and t["sla"] == "SA2" and not t["inoc"] and (r.get("inoc_name") or "").strip():
                    n = _import_ticket(r, key, now, user)
                    t.update(inoc=n["inoc"], topo=n["topo"], ch=n["ch"], create=n["create"], import_key=key)
                    merges.append((pos, t))
                    st["merged"] += 1
                else:
                    st["already"] += 1
                continue
            t = _import_ticket(r, key, now, user)
            new.append(t)
            st["new_closed" if t["status"] == "Closed" else "new_open"] += 1
        if not dry_run and (new or merges):
            if new:
                _state["ws"].append_rows([_ticket_to_row(header, t) for t in new],
                                         value_input_option="RAW", table_range="A1")
            for pos, t in merges:
                _touch(t, user)
                _save(gs, t, pos)
            try:
                lws = _spreadsheet(gs).worksheet(LOG_SHEET)
                stamp = _fmt(now)
                lws.append_rows(
                    [[stamp, t["id"], "import", "IMPORT_SA2_RISK", user["email"], t["import_key"]] for t in new]
                    + [[stamp, t["id"], "import", "IMPORT_SA2_RISK_MERGE", user["email"], t["import_key"]] for _, t in merges],
                    value_input_option="RAW", table_range="A1")
            except Exception:
                log.exception("field-followup import log append failed")
            _load(gs, force=True)
    return st

"""CM Summary Log - Sub Root Cause cluster analysis for planning fiber
optic re-lay work (explicit request: "ใช้ Logic แบบ Operation Temp Point").

Reuses temp_point_improvement.py's clustering engine DIRECTLY (never
duplicated) - build_clusters(), build_summary(), build_province_summary(),
build_improvement_priority(), build_cluster_totals() are all called as-is
from here. This works because the uploaded Excel's "All" sheet uses the
EXACT same column set (Source TT, INC, Site/Cable, CI Name, Severity, FME,
Subject, Activity Owner Group, Sub Root Cause, Type work, Complete Lat Lon,
Month, Region, Province) as the Google Sheet temp_point_improvement.py
reads - confirmed by direct inspection of the uploaded file - so a row
built here in the same dict shape is indistinguishable to that engine from
one built there. Region/Province come as direct columns in this file
(no TRUEOWNERGROUP-style parsing needed, unlike the Google Sheet source).

Loaded via a one-time manual Excel upload (same pattern as
sla_improvement.py's import flow), kept in memory (module-level _store)
until the next import or a server restart - never a live Google Sheet.

Sub Root Cause has ~400+ distinct values in a real export, and Over-SLA-
style clustering runs an OSRM road-alignment check per ambiguous chain
edge - running that across the full ~30k+ row file on every request would
be needlessly slow and hammers the public OSRM server. So clustering only
runs once the person has picked at least one Sub Root Cause (explicit
multi-select request) - with nothing selected, the response carries the
distinct Sub Root Cause list for the filter UI and an awaiting_selection
flag, but no clustering work is done.
"""

import io
import logging
from datetime import datetime

import openpyxl

import temp_point_improvement as tpi

log = logging.getLogger(__name__)

_store = {"rows": None, "warnings": None, "imported_at": None, "filename": None}

SHEET_NAME = "All"

REQUIRED_COLUMNS = [
    "Region", "Province", "Source TT", "INC", "Site/Cable", "CI Name",
    "Severity", "FME", "Subject", "Activity Owner Group", "Sub Root Cause",
    "Type work", "Complete Lat Lon", "Month",
]


class ImportValidationError(Exception):
    pass


def parse_excel(file_bytes):
    """Reads the "All" sheet (explicit request) with openpyxl in
    read-only/data_only mode (values, not formulas - this file has none,
    but stays consistent with how the rest of this codebase reads
    Excel). Raises ImportValidationError with a clear message on a
    missing sheet or missing required column - never silently proceeds
    with a wrong sheet or misaligned columns."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as e:
        raise ImportValidationError(f"เปิดไฟล์ Excel ไม่สำเร็จ: {e}")

    if SHEET_NAME not in wb.sheetnames:
        raise ImportValidationError(f'ไม่พบ Sheet "{SHEET_NAME}" ในไฟล์ (มี Sheet: {", ".join(wb.sheetnames)})')
    ws = wb[SHEET_NAME]

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_raw = next(rows_iter)
    except StopIteration:
        raise ImportValidationError(f'Sheet "{SHEET_NAME}" ไม่มีข้อมูล')
    header = [str(h).strip() if h is not None else "" for h in header_raw]
    header_idx = {h: i for i, h in enumerate(header)}
    missing = [c for c in REQUIRED_COLUMNS if c not in header_idx]
    if missing:
        raise ImportValidationError(f'Sheet "{SHEET_NAME}" ขาด Column ที่จำเป็น: {", ".join(missing)}')

    def get(raw, col):
        idx = header_idx.get(col)
        if idx is None or idx >= len(raw):
            return ""
        v = raw[idx]
        return "" if v is None else str(v).strip()

    rows = []
    invalid_coord_count = 0
    blank_region_count = 0
    seen_point_keys = set()
    duplicate_count = 0

    for raw in rows_iter:
        if raw is None or not any(c is not None and str(c).strip() for c in raw):
            continue  # fully blank row

        region = get(raw, "Region")
        province = get(raw, "Province")
        if not region or not province:
            blank_region_count += 1

        latlon = tpi._parse_complete_lat_lon(get(raw, "Complete Lat Lon"))
        if latlon is None:
            invalid_coord_count += 1

        point_key = (get(raw, "Source TT"), get(raw, "INC"))
        is_duplicate = point_key in seen_point_keys and point_key != ("", "")
        if is_duplicate:
            duplicate_count += 1
        seen_point_keys.add(point_key)

        rows.append({
            "source_tt": get(raw, "Source TT"), "inc": get(raw, "INC"),
            "site_cable": get(raw, "Site/Cable"), "ci_name": get(raw, "CI Name"),
            "severity": get(raw, "Severity"), "fme": get(raw, "FME"),
            "subject": get(raw, "Subject"), "activity_owner_group": get(raw, "Activity Owner Group"),
            "sub_root_cause": get(raw, "Sub Root Cause"), "type_work": get(raw, "Type work"),
            "complete_lat_lon_raw": get(raw, "Complete Lat Lon"), "month": get(raw, "Month"),
            "region": region or None, "province": province or None,
            "latitude": latlon[0] if latlon else None, "longitude": latlon[1] if latlon else None,
            "invalid_coord": latlon is None, "is_duplicate": is_duplicate,
        })

    warnings = []
    if blank_region_count:
        warnings.append(f"{blank_region_count} แถวไม่มี Region/Province")
    if invalid_coord_count:
        warnings.append(f"{invalid_coord_count} แถวมีพิกัดไม่ถูกต้อง/ไม่มีพิกัด (Invalid/Unmapped)")
    if duplicate_count:
        warnings.append(f"พบ {duplicate_count} แถวซ้ำ (Source TT + INC ซ้ำกัน)")

    return rows, warnings


def import_excel(file_bytes, filename=None):
    rows, warnings = parse_excel(file_bytes)
    _store["rows"] = rows
    _store["warnings"] = warnings
    _store["imported_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    _store["filename"] = filename
    return {"row_count": len(rows), "warnings": warnings}


def get_rows():
    return _store["rows"]


def distinct_sub_root_causes(rows):
    """Every distinct non-blank Sub Root Cause present, with how many
    rows carry it - lets the filter UI show counts and sort by
    frequency, useful given a real export has 400+ distinct values."""
    counts = {}
    for r in rows:
        c = r["sub_root_cause"]
        if c:
            counts[c] = counts.get(c, 0) + 1
    return sorted(({"label": c, "count": n} for c, n in counts.items()), key=lambda x: -x["count"])


def build_cm_summary_response(sub_root_causes=None):
    """sub_root_causes: list of Sub Root Cause values to keep, or falsy
    for "nothing selected yet" - returns an awaiting_selection response
    with just the filter options in that case (see module docstring for
    why clustering doesn't run unfiltered by default)."""
    rows = _store["rows"]
    if rows is None:
        return None

    all_causes = distinct_sub_root_causes(rows)
    base = {
        "imported_at": _store["imported_at"], "filename": _store["filename"],
        "import_warnings": _store["warnings"], "row_count": len(rows),
        "sub_root_causes": all_causes,
    }

    if not sub_root_causes:
        base["awaiting_selection"] = True
        return base

    wanted = set(sub_root_causes)
    filtered = [r for r in rows if r["sub_root_cause"] in wanted]

    clustered, invalid = tpi.build_clusters(filtered)
    summary = tpi.build_summary(clustered, invalid)
    province_summary = tpi.build_province_summary(clustered)
    province_priority = tpi.build_improvement_priority(province_summary)
    cluster_totals = tpi.build_cluster_totals(clustered)

    points_out = [{
        "cluster_id": p["cluster_id"], "is_cluster": p["is_cluster"],
        "source_tt": p["source_tt"], "inc": p["inc"], "site_cable": p["site_cable"],
        "ci_name": p["ci_name"], "severity": p["severity"], "fme": p["fme"], "subject": p["subject"],
        "activity_owner_group": p["activity_owner_group"], "sub_root_cause": p["sub_root_cause"],
        "type_work": p["type_work"], "month": p["month"],
        "region": p["region"], "province": p["province"],
        "latitude": p["latitude"], "longitude": p["longitude"],
        "distance_to_previous_m": p["distance_to_previous_m"], "cluster_distance_m": p["cluster_distance_m"],
    } for p in clustered]
    invalid_out = [{
        "source_tt": p["source_tt"], "inc": p["inc"], "site_cable": p["site_cable"],
        "ci_name": p["ci_name"], "severity": p["severity"], "province": p["province"],
        "sub_root_cause": p["sub_root_cause"], "complete_lat_lon_raw": p["complete_lat_lon_raw"],
    } for p in invalid]

    base.update({
        "awaiting_selection": False,
        "selected_sub_root_causes": sorted(wanted),
        "filtered_row_count": len(filtered),
        "summary": summary, "province_summary": province_priority,
        "cluster_totals": cluster_totals, "points": points_out,
        "invalid_points": invalid_out,
        "detail_table_columns": tpi.DETAIL_TABLE_COLUMNS,
    })
    return base


def build_export_rows(sub_root_causes):
    rows = _store["rows"]
    if rows is None or not sub_root_causes:
        return []
    wanted = set(sub_root_causes)
    filtered = [r for r in rows if r["sub_root_cause"] in wanted]
    clustered, _ = tpi.build_clusters(filtered)
    return tpi.build_export_rows(clustered)

"""Historical Closed Ticket - lookup ticket history by CI Name (explicit
request). Manually imported Excel (single sheet, first sheet used
regardless of its name - the real exports seen so far are single-sheet,
name not guaranteed to stay the same), kept in memory (module-level
_store) like sla_improvement.py's import flow - not a live Google Sheet.

Search logic (explicit request, confirmed choice): exact match against
the CI name column first; if that finds nothing, fall back to a
substring match against SUBJECT (case-insensitive, matches anywhere in
the text) - since SUBJECT often embeds the CI name as a prefix (e.g. CI
name "CMI8506" inside SUBJECT "CMI8506L_6NB03 (A71+ORIGIN): ERIC Link
Failure") but isn't always identical to it, so a ticket can be relevant
even when the CI name field itself is blank or doesn't match exactly.

Column set (explicit request, second revision - source file's CI-name
column is now literally named "CINAME", not "CI_Name" as in the first
file this tab was built against; CLOSEDTIME and REMEDY no longer exist
in the source at all, replaced by RESTORATIONDATE + STATUS). CI_NAME_COL
tries "CINAME" first and falls back to "CI_Name" so either file schema
still imports without a code change.
"""

import io
import logging
from datetime import datetime

import openpyxl

log = logging.getLogger(__name__)

_store = {"rows": None, "warnings": None, "imported_at": None, "filename": None}

CI_NAME_COL_CANDIDATES = ["CINAME", "CI_Name"]

REQUIRED_COLUMNS = [
    "TICKETID", "SUBJECT", "CREATIONDATE", "RESTORATIONDATE", "STATUS",
    "TRUESEVERITY_DESC", "NODENAME", "CLASSIFICATION", "TRUEOWNERGROUP",
    "PROBLEM", "SUB_CAUSE", "DISTRICT_EN",
]

# Column display order (explicit request, second revision): CREATIONDATE,
# RESTORATIONDATE, STATUS inserted right after it, TRUESEVERITY_DESC,
# then the rest unchanged - REMEDY dropped entirely (no longer in the
# source file).
RESULT_COLUMN_ORDER = [
    "CREATIONDATE", "RESTORATIONDATE", "STATUS", "TRUESEVERITY_DESC",
    "TICKETID", "SUBJECT", "NODENAME", "CLASSIFICATION", "TRUEOWNERGROUP",
    "PROBLEM", "SUB_CAUSE", "DISTRICT_EN",
]

# Fields to summarize as pie charts once a search has results (explicit request).
BREAKDOWN_FIELDS = ["STATUS", "TRUESEVERITY_DESC", "PROBLEM", "SUB_CAUSE"]


class ImportValidationError(Exception):
    pass


def _parse_dt(v):
    """Cell values come back as datetime objects already when the sheet
    stores them as real dates (openpyxl with data_only=True) - falls back
    to string parsing only if the cell was stored as text."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v
    s = str(v).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%d/%m/%Y %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def parse_excel(file_bytes):
    """Reads the FIRST sheet regardless of its name. Raises
    ImportValidationError with a clear message on a missing required
    column, or if NEITHER CI-name column candidate is present - never
    silently proceeds with misaligned columns."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    except Exception as e:
        raise ImportValidationError(f"เปิดไฟล์ Excel ไม่สำเร็จ: {e}")

    ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_raw = next(rows_iter)
    except StopIteration:
        raise ImportValidationError("ไฟล์ไม่มีข้อมูล")
    header = [str(h).strip() if h is not None else "" for h in header_raw]
    header_idx = {h: i for i, h in enumerate(header)}

    missing = [c for c in REQUIRED_COLUMNS if c not in header_idx]
    if missing:
        raise ImportValidationError(f"ไฟล์ขาด Column ที่จำเป็น: {', '.join(missing)}")
    ci_name_col = next((c for c in CI_NAME_COL_CANDIDATES if c in header_idx), None)
    if ci_name_col is None:
        raise ImportValidationError(f"ไฟล์ขาด Column ชื่อ CI Name (ลองหาแล้ว: {', '.join(CI_NAME_COL_CANDIDATES)})")

    def get(raw, col):
        idx = header_idx.get(col)
        if idx is None or idx >= len(raw):
            return None
        return raw[idx]

    def get_str(raw, col):
        v = get(raw, col)
        return "" if v is None else str(v).strip()

    rows = []
    blank_ci_name_count = 0
    for raw in rows_iter:
        if raw is None or not any(c is not None and str(c).strip() for c in raw):
            continue  # fully blank row

        ci_name = get_str(raw, ci_name_col)
        if not ci_name:
            blank_ci_name_count += 1

        rows.append({
            "TICKETID": get_str(raw, "TICKETID"),
            "SUBJECT": get_str(raw, "SUBJECT"),
            "CREATIONDATE": _parse_dt(get(raw, "CREATIONDATE")),
            "RESTORATIONDATE": _parse_dt(get(raw, "RESTORATIONDATE")),
            "STATUS": get_str(raw, "STATUS"),
            "TRUESEVERITY_DESC": get_str(raw, "TRUESEVERITY_DESC"),
            "NODENAME": get_str(raw, "NODENAME"),
            "CLASSIFICATION": get_str(raw, "CLASSIFICATION"),
            "TRUEOWNERGROUP": get_str(raw, "TRUEOWNERGROUP"),
            "PROBLEM": get_str(raw, "PROBLEM"),
            "SUB_CAUSE": get_str(raw, "SUB_CAUSE"),
            "DISTRICT_EN": get_str(raw, "DISTRICT_EN"),
            "CI_NAME": ci_name,
        })

    warnings = []
    if blank_ci_name_count:
        warnings.append(f"{blank_ci_name_count} แถวไม่มีค่า CI Name (ยังค้นหาเจอได้ผ่าน SUBJECT)")

    return rows, warnings


def import_excel(file_bytes, filename=None):
    rows, warnings = parse_excel(file_bytes)
    _store["rows"] = rows
    _store["warnings"] = warnings
    _store["imported_at"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    _store["filename"] = filename
    return {"row_count": len(rows), "warnings": warnings}


def get_rows():
    return _store["rows"]


def distinct_severities(rows):
    return sorted({r["TRUESEVERITY_DESC"] for r in rows if r["TRUESEVERITY_DESC"]})


def distinct_statuses(rows):
    return sorted({r["STATUS"] for r in rows if r["STATUS"]})


def _fmt_dt(dt):
    return dt.strftime("%Y-%m-%d %H:%M:%S") if dt else None


def _build_breakdowns(matched_rows):
    """One {label, count, pct} list per BREAKDOWN_FIELDS field (explicit
    request - feeds the 4 pie charts shown after a search). Blank values
    are grouped under "(ไม่ระบุ)" rather than dropped, so the percentages
    still add up to 100% of the matched tickets."""
    out = {}
    total = len(matched_rows)
    for field in BREAKDOWN_FIELDS:
        counts = {}
        for r in matched_rows:
            label = r[field] or "(ไม่ระบุ)"
            counts[label] = counts.get(label, 0) + 1
        rows_out = [
            {"label": label, "count": n, "pct": round(n / total * 100, 1) if total else 0}
            for label, n in counts.items()
        ]
        rows_out.sort(key=lambda x: -x["count"])
        out[field] = rows_out
    return out


def search(ci_name_query, severities=None, statuses=None):
    """Returns (results, match_mode, breakdowns, error). match_mode is
    "exact" (found via the CI-name column), "subject" (fell back to
    SUBJECT substring), or None (no query typed / nothing found / no data
    yet). Results are ALWAYS sorted by CREATIONDATE oldest -> newest
    (explicit request), regardless of which match mode found them.
    breakdowns is None when there are no results to summarize."""
    rows = _store["rows"]
    if rows is None:
        return None, None, None, "ยังไม่มีข้อมูล กรุณา Import ไฟล์ก่อน"

    query = (ci_name_query or "").strip()
    if not query:
        return [], None, None, None

    query_lower = query.lower()
    matched = [r for r in rows if r["CI_NAME"].lower() == query_lower]
    match_mode = "exact"
    if not matched:
        matched = [r for r in rows if query_lower in r["SUBJECT"].lower()]
        match_mode = "subject"

    if severities:
        wanted = set(severities)
        matched = [r for r in matched if r["TRUESEVERITY_DESC"] in wanted]
    if statuses:
        wanted_status = set(statuses)
        matched = [r for r in matched if r["STATUS"] in wanted_status]

    matched.sort(key=lambda r: r["CREATIONDATE"] or datetime.min)

    results = [{
        "CREATIONDATE": _fmt_dt(r["CREATIONDATE"]),
        "RESTORATIONDATE": _fmt_dt(r["RESTORATIONDATE"]),
        "STATUS": r["STATUS"],
        "TRUESEVERITY_DESC": r["TRUESEVERITY_DESC"],
        "TICKETID": r["TICKETID"],
        "SUBJECT": r["SUBJECT"],
        "NODENAME": r["NODENAME"],
        "CLASSIFICATION": r["CLASSIFICATION"],
        "TRUEOWNERGROUP": r["TRUEOWNERGROUP"],
        "PROBLEM": r["PROBLEM"],
        "SUB_CAUSE": r["SUB_CAUSE"],
        "DISTRICT_EN": r["DISTRICT_EN"],
        "CI_NAME": r["CI_NAME"],
    } for r in matched]

    breakdowns = _build_breakdowns(matched) if matched else None

    return results, (match_mode if results else None), breakdowns, None
